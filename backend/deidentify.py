"""De-identify a document's text and restore it later from a key.

Everything here works in memory: nothing is written to disk or the database.
Placeholders look like [PERSON_1] or [UCI_2]; the key maps each placeholder
back to the original text and is handed to the user, never kept by the server.

What is hidden follows the redaction practice of published RPD decisions
(names of claimants and other people, identity numbers, contact details,
specific dates, ages), with two deliberate exceptions chosen by the owner:
countries and employment details stay in the text. IRB and court file numbers
are hidden too: published decisions keep them, but they link straight to the
person.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
import re
from typing import Any, Iterable

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
KEY_FORMAT = "ilit-deidentify-key"
KEY_VERSION = 1

# Categories in the order they win when two matches overlap at the same length.
CATEGORY_LABELS = {
	"PERSON": "People",
	"DETAIL": "Other details you listed",
	"EMAIL": "Email addresses",
	"PHONE": "Phone numbers",
	"SIN": "Social insurance numbers",
	"UCI": "UCI / client ID numbers",
	"APPLICATION": "IRCC application numbers",
	"IRB_FILE": "IRB file numbers",
	"COURT_FILE": "Federal Court file numbers",
	"PASSPORT": "Passport numbers",
	"ID": "Other ID / file numbers",
	"ADDRESS": "Street addresses",
	"POSTAL": "Postal codes",
	"DATE": "Specific dates",
	"AGE": "Ages",
}
CATEGORY_PRIORITY = {name: index for index, name in enumerate(CATEGORY_LABELS)}

_MONTHS = (
	r"(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|June?|July?|Aug(?:ust)?|"
	r"Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)\.?"
)
_STREET_TYPES = (
	r"(?:Street|St|Avenue|Ave|Road|Rd|Boulevard|Blvd|Drive|Dr|Crescent|Cres|Court|Ct|Lane|Ln|Way|"
	r"Place|Pl|Terrace|Terr|Circle|Cir|Trail|Trl|Parkway|Pkwy|Highway|Hwy|Square|Sq|Gate|Grove|"
	r"Heights|Hts|Row|Close)"
)
_ID_KEYWORDS = (
	r"(?:file|client|application|reference|ref|case|document|permit|licen[cs]e|certificate|"
	r"travel\s+document|national\s+id(?:entity)?(?:\s+card)?|identity\s+card|id|"
	r"health\s+card|driver'?s\s+licen[cs]e|account)"
)

PATTERNS: list[tuple[str, re.Pattern[str]]] = [
	("EMAIL", re.compile(r"\b[\w.+-]+@[\w-]+(?:\.[\w-]+)+\b")),
	(
		"PHONE",
		re.compile(
			r"(?<![\w-])(?:\+?1[\s.-]?)?\(?\d{3}\)?[\s.-]\d{3}[\s.-]\d{4}(?:\s*(?:ext\.?|x)\s*\d{1,5})?(?![\w-])"
			r"|(?<![\w-])\+\d{1,3}(?:[\s.-]?\(?\d{1,4}\)?){2,5}(?![\w-])"
		),
	),
	("IRB_FILE", re.compile(r"\b[A-Z]{2}\d-\d{5}\b")),
	("COURT_FILE", re.compile(r"\bIMM-\d{1,6}-\d{2}\b", re.IGNORECASE)),
	("APPLICATION", re.compile(r"\b[BEFSVWX]\d{9}\b")),
	("UCI", re.compile(r"(?<![\w-])(?:\d{2}-\d{4}-\d{4}|\d{4}-\d{4})(?![\w-])")),
	("POSTAL", re.compile(r"\b[ABCEGHJ-NPRSTVXY]\d[ABCEGHJ-NPRSTV-Z][ -]?\d[ABCEGHJ-NPRSTV-Z]\d\b")),
	(
		"ADDRESS",
		re.compile(
			r"(?:\b(?:Apt|Apartment|Unit|Suite)\.?\s*#?\s*\w{1,6},?\s+)?"
			r"\b\d{1,6}[A-Za-z]?(?:-\d{1,6})?\s+(?:[A-Z][\w'.-]*\s+){1,4}" + _STREET_TYPES + r"\b\.?"
			r"(?:\s+(?:North|South|East|West|N|S|E|W)\b\.?)?"
			r"(?:,?\s+(?:Apt|Apartment|Unit|Suite)\.?\s*#?\s*\w{1,6})?"
		),
	),
	(
		"DATE",
		re.compile(
			r"\b(?:\d{1,2}(?:st|nd|rd|th)?\s+" + _MONTHS + r",?\s+\d{4}"
			r"|" + _MONTHS + r"\s+\d{1,2}(?:st|nd|rd|th)?,?\s+\d{4}"
			r"|(?:19|20)\d{2}[-/.](?:0?[1-9]|1[0-2])[-/.](?:0?[1-9]|[12]\d|3[01])"
			r"|(?:0?[1-9]|[12]\d|3[01])[-/.](?:0?[1-9]|1[0-2])[-/.](?:19|20)\d{2}"
			r"|(?:0?[1-9]|1[0-2])[-/.](?:0?[1-9]|[12]\d|3[01])[-/.](?:19|20)\d{2})\b",
			re.IGNORECASE,
		),
	),
	(
		"AGE",
		re.compile(r"\b\d{1,3}[- ]years?[- ]old\b|\b(?:aged?|age of)\s+\d{1,3}\b", re.IGNORECASE),
	),
]

# Numbers that only count when a keyword sits just before them.
_SIN_CONTEXT = re.compile(
	r"\b(?:SIN|social\s+insurance\s+(?:number|no\.?|#))\s*(?:number|no\.?|#)?\s*[:#]?\s*(\d{3}[\s-]?\d{3}[\s-]?\d{3})\b",
	re.IGNORECASE,
)
_SIN_FORMATTED = re.compile(r"(?<![\w-])\d{3}[ -]\d{3}[ -]\d{3}(?![\w-])")
_PASSPORT_CONTEXT = re.compile(
	r"\bpassport\s*(?:number|no\.?|#)?\s*[:#]?\s*([A-Z]{0,3}\d[A-Z0-9]{4,11})\b",
	re.IGNORECASE,
)
_ID_CONTEXT = re.compile(
	r"\b" + _ID_KEYWORDS + r"\s*(?:number|no\.?|num\.?|#)\s*[:#]?\s*([A-Z0-9][A-Z0-9-]{3,24})\b",
	re.IGNORECASE,
)
_UCI_CONTEXT = re.compile(r"\b(?:UCI|client\s+ID)\s*(?:number|no\.?|#)?\s*[:#]?\s*(\d{8}|\d{10})\b", re.IGNORECASE)

_PLACEHOLDER = re.compile(r"\[\s*([A-Za-z]+(?:[ _-][A-Za-z]+)*)[ _-]*(\d+)((?:[ _-][A-Za-z][A-Za-z0-9]*){0,2})\s*\]")
_HONORIFICS = {"mr", "mrs", "ms", "miss", "mx", "dr", "sr", "jr", "mme", "m", "mlle"}


@dataclass(frozen=True)
class Span:
	start: int
	end: int
	category: str
	group: str  # normalised identity: all spans with the same group share a placeholder number
	variant: str = ""  # suffix for parts of a name, e.g. SURNAME


def _luhn_ok(digits: str) -> bool:
	total = 0
	for index, char in enumerate(reversed(digits)):
		value = int(char)
		if index % 2 == 1:
			value *= 2
			if value > 9:
				value -= 9
		total += value
	return total % 10 == 0


def _digits(value: str) -> str:
	return re.sub(r"\D", "", value)


def _norm(value: str) -> str:
	return re.sub(r"\s+", " ", value).strip().casefold()


def _looks_like_year_range(value: str) -> bool:
	parts = value.split("-")
	return len(parts) == 2 and all(re.fullmatch(r"(?:19|20)\d{2}", part) for part in parts)


def _name_regex(value: str) -> str:
	return r"(?<![\w])" + r"\s+".join(re.escape(part) for part in value.split()) + r"(?![\w])"


def _person_spans(text: str, names: Iterable[str]) -> list[Span]:
	spans: list[Span] = []
	for person_index, name in enumerate(names):
		parts = [part.strip(".,") for part in name.split()]
		parts = [part for part in parts if part and part.casefold() not in _HONORIFICS]
		if not parts:
			continue
		group = f"person:{person_index}"
		forms: list[tuple[str, str]] = [(" ".join(parts), "")]
		if len(parts) > 1:
			forms.append((f"{parts[-1]}, {' '.join(parts[:-1])}", ""))
			forms.append((f"{parts[-1]} {' '.join(parts[:-1])}", ""))
			forms.append((parts[-1], "SURNAME"))
			forms.append((parts[0], "GIVEN"))
			for middle_index, middle in enumerate(parts[1:-1], start=2):
				forms.append((middle, f"NAME{middle_index}"))
		for form, variant in forms:
			if len(form) < 2:
				continue
			for match in re.finditer(_name_regex(form), text, re.IGNORECASE):
				spans.append(Span(match.start(), match.end(), "PERSON", group, variant))
	return spans


def _detail_spans(text: str, details: Iterable[str]) -> list[Span]:
	spans: list[Span] = []
	for detail in details:
		if len(detail) < 2:
			continue
		for match in re.finditer(_name_regex(detail), text, re.IGNORECASE):
			spans.append(Span(match.start(), match.end(), "DETAIL", f"detail:{_norm(detail)}"))
	return spans


def _pattern_spans(text: str, categories: set[str]) -> list[Span]:
	spans: list[Span] = []
	for category, pattern in PATTERNS:
		if category not in categories:
			continue
		for match in pattern.finditer(text):
			value = match.group(0)
			if category == "UCI" and _looks_like_year_range(value):
				continue
			if category in {"UCI", "PHONE", "APPLICATION"}:
				group = _digits(value) or _norm(value)
			else:
				group = _norm(value)
			spans.append(Span(match.start(), match.end(), category, f"{category}:{group}"))

	contextual = [
		("SIN", _SIN_CONTEXT),
		("PASSPORT", _PASSPORT_CONTEXT),
		("UCI", _UCI_CONTEXT),
		("ID", _ID_CONTEXT),
	]
	for category, pattern in contextual:
		if category not in categories:
			continue
		for match in pattern.finditer(text):
			value = match.group(1)
			if category == "ID" and not re.search(r"\d", value):
				continue
			group = _digits(value) if category in {"SIN", "UCI"} else _norm(value)
			spans.append(Span(match.start(1), match.end(1), category, f"{category}:{group}"))
	if "SIN" in categories:
		for match in _SIN_FORMATTED.finditer(text):
			if _luhn_ok(_digits(match.group(0))):
				spans.append(Span(match.start(), match.end(), "SIN", f"SIN:{_digits(match.group(0))}"))
	return spans


def _resolve_overlaps(spans: list[Span]) -> list[Span]:
	ordered = sorted(spans, key=lambda s: (-(s.end - s.start), CATEGORY_PRIORITY[s.category], s.start))
	taken: list[Span] = []
	for span in ordered:
		if any(span.start < kept.end and kept.start < span.end for kept in taken):
			continue
		taken.append(span)
	return sorted(taken, key=lambda s: s.start)


def _numeric_groups(group: str) -> str:
	# Give UCI/SIN/ID numbers that share digits the same placeholder number.
	category, _, value = group.partition(":")
	if category in {"UCI", "SIN", "ID", "PASSPORT", "APPLICATION"}:
		digits = _digits(value)
		return f"NUM:{digits}" if digits else group
	return group


def deidentify_text(
	text: str,
	names: Iterable[str] = (),
	details: Iterable[str] = (),
	categories: Iterable[str] | None = None,
	source_name: str = "",
) -> dict[str, Any]:
	"""Replace identifying details with placeholders. Returns text, key, summary and warnings."""
	names = [name for name in names if name.strip()]
	details = [detail for detail in details if detail.strip()]
	enabled = set(categories) if categories is not None else set(CATEGORY_LABELS)
	enabled |= {"PERSON", "DETAIL"}

	spans = _person_spans(text, names) + _detail_spans(text, details) + _pattern_spans(text, enabled)
	spans = _resolve_overlaps(spans)

	numbers: dict[str, dict[str, int]] = {}
	group_numbers: dict[str, int] = {}
	surface_forms: dict[str, Counter[str]] = {}
	pieces: list[str] = []
	cursor = 0
	for span in spans:
		identity = _numeric_groups(span.group)
		per_category = numbers.setdefault(span.category, {})
		if identity not in per_category:
			per_category[identity] = len(per_category) + 1
		group_numbers[span.group] = per_category[identity]
		surface = text[span.start : span.end]
		variants = [span.variant] if span.variant else []
		if span.category in {"PERSON", "DETAIL"} and surface.isupper() and len(surface) > 1:
			# Keep ALL-CAPS spellings (common in court headings) separate so they restore as written.
			variants.append("CAPS")
		placeholder = "_".join([span.category, str(per_category[identity]), *variants])
		surface_forms.setdefault(placeholder, Counter())[surface] += 1
		pieces.append(text[cursor : span.start])
		pieces.append(f"[{placeholder}]")
		cursor = span.end
	pieces.append(text[cursor:])
	redacted = "".join(pieces)

	entries = {placeholder: forms.most_common(1)[0][0] for placeholder, forms in surface_forms.items()}
	counts: Counter[str] = Counter()
	for span in spans:
		counts[span.category] += 1

	warnings: list[str] = []
	folded = redacted.casefold()
	for original in sorted({value for forms in surface_forms.values() for value in forms}):
		if len(original) >= 3 and original.casefold() in folded:
			warnings.append(f"“{original}” still appears in the de-identified text.")
	missing = [name for index, name in enumerate(names) if not any(s.group == f"person:{index}" for s in spans)]
	for name in missing:
		warnings.append(f"The name “{name}” was not found in the document. Check the spelling.")

	return {
		"text": redacted,
		"key": {
			"format": KEY_FORMAT,
			"version": KEY_VERSION,
			"created": datetime.now(timezone.utc).isoformat(timespec="seconds"),
			"source": source_name,
			"entries": entries,
		},
		"summary": [
			{"category": category, "label": CATEGORY_LABELS[category], "count": counts[category]}
			for category in CATEGORY_LABELS
			if counts[category]
		],
		"replacements": sum(counts.values()),
		"warnings": warnings,
	}


def _canonical_placeholder(match: re.Match[str]) -> str:
	label = re.sub(r"[ -]", "_", match.group(1)).upper()
	suffix = re.sub(r"[ -]", "_", match.group(3) or "").upper()
	if suffix and not suffix.startswith("_"):
		suffix = "_" + suffix
	return f"{label}_{match.group(2)}{suffix}"


def validate_key(key: Any) -> dict[str, str]:
	if not isinstance(key, dict) or key.get("format") != KEY_FORMAT:
		raise ValueError("This is not an iLit de-identification key file.")
	entries = key.get("entries")
	if not isinstance(entries, dict) or not all(isinstance(k, str) and isinstance(v, str) for k, v in entries.items()):
		raise ValueError("The key file is damaged: its entries are missing or malformed.")
	return {k.upper(): v for k, v in entries.items()}


def reidentify_text(text: str, key: Any) -> dict[str, Any]:
	"""Put the original details back in place of placeholders."""
	entries = validate_key(key)
	used: Counter[str] = Counter()
	unknown: list[str] = []

	def replace(match: re.Match[str]) -> str:
		placeholder = _canonical_placeholder(match)
		if placeholder in entries:
			used[placeholder] += 1
			return entries[placeholder]
		if any(placeholder.startswith(category + "_") for category in CATEGORY_LABELS):
			unknown.append(match.group(0))
		return match.group(0)

	restored = _PLACEHOLDER.sub(replace, text)
	warnings: list[str] = []
	if unknown:
		listed = ", ".join(sorted(set(unknown))[:10])
		warnings.append(f"These placeholders are not in the key and were left as they are: {listed}.")
	not_found = [placeholder for placeholder in entries if not used[placeholder]]
	return {
		"text": restored,
		"restored": sum(used.values()),
		"unused_placeholders": [f"[{placeholder}]" for placeholder in sorted(not_found)],
		"warnings": warnings,
	}


# ---- Reading uploaded files -------------------------------------------------


def _docx_block_text(element: Any) -> Iterable[str]:
	if element.tag == qn("w:p"):
		yield "".join(node.text or "" for node in element.iter(qn("w:t")))
	elif element.tag == qn("w:tbl"):
		for row in element.iter(qn("w:tr")):
			cells = []
			for cell in row.iter(qn("w:tc")):
				cells.append(" ".join("".join(t.text or "" for t in p.iter(qn("w:t"))) for p in cell.iter(qn("w:p"))).strip())
			yield " | ".join(cells)


def _text_from_docx(content: bytes) -> str:
	document = Document(BytesIO(content))
	blocks: list[str] = []
	seen_parts: set[int] = set()
	for section in document.sections:
		for part in (section.header, section.first_page_header):
			if part.is_linked_to_previous or id(part._element) in seen_parts:
				continue
			seen_parts.add(id(part._element))
			for element in part._element.iterchildren():
				blocks.extend(_docx_block_text(element))
	for element in document.element.body.iterchildren():
		blocks.extend(_docx_block_text(element))
	for section in document.sections:
		for part in (section.footer, section.first_page_footer):
			if part.is_linked_to_previous or id(part._element) in seen_parts:
				continue
			seen_parts.add(id(part._element))
			for element in part._element.iterchildren():
				blocks.extend(_docx_block_text(element))
	return "\n\n".join(block for block in blocks if block.strip())


def _text_from_pdf(content: bytes) -> str:
	reader = PdfReader(BytesIO(content))
	pages = [(page.extract_text() or "").strip() for page in reader.pages]
	if pages and sum(len(page) for page in pages) < 40 * len(pages):
		raise ValueError(
			"This PDF looks scanned (it has little or no selectable text). Scanned PDFs are not supported yet."
		)
	return "\n\n".join(pages)


def text_from_upload(filename: str | None, content: bytes) -> str:
	name = (filename or "").lower()
	if not content:
		raise ValueError("The uploaded file is empty.")
	if len(content) > MAX_UPLOAD_BYTES:
		raise ValueError("The uploaded file is over the 10 MB limit.")
	if name.endswith(".docx"):
		return _text_from_docx(content)
	if name.endswith(".pdf"):
		return _text_from_pdf(content)
	if name.endswith((".txt", ".md")):
		return content.decode("utf-8-sig", errors="replace")
	raise ValueError("Only .docx, .pdf and .txt files are supported.")


def text_to_docx(text: str) -> bytes:
	document = Document()
	for block in re.split(r"\n\s*\n", text):
		document.add_paragraph(block.strip("\n"))
	buffer = BytesIO()
	document.save(buffer)
	return buffer.getvalue()
