"""DOCX serialization for the source-linked legal issue brief."""

from __future__ import annotations

from datetime import datetime, timezone
from typing import Any

from docx.shared import Pt

from .docx_export import add_docx_hyperlink, new_docx_document, serialize_docx
from .issue_brief_analytics import minister_win_rate_label


MAX_PRINT_DECISIONS = 12


def issue_brief_docx(brief: dict[str, Any]) -> bytes:
	"""Render the issue-brief page facts and traceable links into Word format."""
	tag = str(brief.get("tag") or "")
	decision_count = int(brief.get("decision_count") or 0)
	semantics = brief.get("semantics") or {}
	title = f"Legal issue brief: {tag or '(empty tag)'}"
	document = new_docx_document(title)
	document.paragraphs[0].style = "Heading 1"
	document.add_paragraph(
		f"{decision_count} tagged decisions · exact active-taxonomy tag match"
	)
	if not tag:
		document.add_paragraph(
			"Enter a tag in category:value form, for example issue:procedural_fairness, "
			"then reload /issue-brief-ui?tag=…."
		)
	if not decision_count and tag:
		document.add_paragraph("No decisions are tagged with this issue.")

	if decision_count:
		document.add_heading("Decisions by year and outcome", level=2)
		years_table = document.add_table(rows=1, cols=4)
		years_table.style = "Table Grid"
		for cell, label in zip(
			years_table.rows[0].cells,
			(
				"Year / outcome",
				"Decisions",
				"Unclassified / outcome share",
				"Minister win rate",
			),
		):
			cell.text = label
		for row in brief.get("years", []):
			year_cells = years_table.add_row().cells
			year_cells[0].text = str(
				row.get("year") if row.get("year") is not None else "Unknown"
			)
			year_cells[1].text = str(int(row.get("decision_count") or 0))
			year_cells[2].text = (
				f"Unclassified: {int(row.get('unclassified_count') or 0)}"
			)
			year_cells[3].text = minister_win_rate_label(row.get("minister_win_rate"))
			for outcome in row.get("outcome_splits", []):
				outcome_cells = years_table.add_row().cells
				outcome_cells[0].text = str(outcome.get("outcome") or "unclassified")
				outcome_cells[1].text = str(int(outcome.get("count") or 0))
				outcome_cells[2].text = (
					f"{float(outcome.get('percentage') or 0):.1f}% "
					f"(unclassified {int(outcome.get('unclassified_count') or 0)}; "
					f"denominator {int(outcome.get('denominator') or 0)})"
				)
		document.add_paragraph(
			"Outcome percentages use all tagged decisions in the year. Minister "
			"win rates use classified government outcomes only; total n and "
			"unclassified counts are shown for each year."
		)

		document.add_heading("Courts", level=2)
		courts_table = document.add_table(rows=1, cols=2)
		courts_table.style = "Table Grid"
		courts_table.rows[0].cells[0].text = "Court"
		courts_table.rows[0].cells[1].text = "Decisions"
		for row in brief.get("courts", []):
			cells = courts_table.add_row().cells
			cells[0].text = str(row.get("court") or "Unspecified")
			cells[1].text = str(int(row.get("decision_count") or 0))

		document.add_heading("Top cited authorities", level=2)
		authorities = brief.get("top_authorities", [])
		if authorities:
			authority_table = document.add_table(rows=1, cols=3)
			authority_table.style = "Table Grid"
			for cell, label in zip(
				authority_table.rows[0].cells,
				("Authority", "Citation occurrences", "Tagged citing decisions"),
			):
				cell.text = label
			for row in authorities:
				cells = authority_table.add_row().cells
				paragraph = cells[0].paragraphs[0]
				label = str(row.get("citation") or row.get("title") or "Linked authority")
				url = str(row.get("url") or "")
				if url:
					add_docx_hyperlink(paragraph, label, url)
				else:
					paragraph.add_run(label)
				cells[1].text = str(int(row.get("citation_occurrences") or 0))
				cells[2].text = str(int(row.get("citing_decisions") or 0))
		else:
			document.add_paragraph("No resolved case citations found.")

		document.add_heading("Tagged decisions", level=2)
		decisions = brief.get("decisions", [])
		for row in decisions[:MAX_PRINT_DECISIONS]:
			paragraph = document.add_paragraph(style="List Bullet")
			label = str(row.get("citation") or row.get("title") or "Decision")
			url = str(row.get("url") or "")
			if url:
				add_docx_hyperlink(paragraph, label, url)
			else:
				paragraph.add_run(label)
			paragraph.add_run(f" — {row.get('court') or 'Unspecified'}")
			if row.get("outcome"):
				paragraph.add_run(f" · {row['outcome']}")
		if not decisions:
			document.add_paragraph("No linked decisions.")
		if len(decisions) or decision_count:
			shown_count = min(len(decisions), MAX_PRINT_DECISIONS)
			document.add_paragraph(
				f"Showing {shown_count} of {decision_count} decisions; "
				"see JSON brief for the complete list."
			)

	document.add_paragraph(f"Outcome source: {semantics.get('outcomes') or ''}")
	document.add_paragraph(f"Citation scope: {semantics.get('citations') or ''}")
	if semantics.get("tag_matching"):
		document.add_paragraph(f"Tag matching: {semantics['tag_matching']}")
	if semantics.get("minister_outcomes"):
		document.add_paragraph(f"Minister outcomes: {semantics['minister_outcomes']}")

	footer = document.sections[0].footer.paragraphs[0]
	footer.text = (
		"Generated from iLit data on "
		+ datetime.now(timezone.utc).date().isoformat()
		+ "; descriptive statistics, see denominators"
	)
	for run in footer.runs:
		run.font.size = Pt(8)
	return serialize_docx(document)
