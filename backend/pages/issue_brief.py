"""Printable, source-linked legal issue brief page."""

from __future__ import annotations

from datetime import datetime, timezone
from html import escape
from io import BytesIO
from typing import Any

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from ..deidentify import text_to_docx


MAX_PRINT_DECISIONS = 12


def _add_hyperlink(paragraph: Any, text: str, url: str) -> None:
	"""Add a clickable link to a Word paragraph."""
	relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
	hyperlink = OxmlElement("w:hyperlink")
	hyperlink.set(qn("r:id"), relationship_id)
	run = OxmlElement("w:r")
	properties = OxmlElement("w:rPr")
	color = OxmlElement("w:color")
	color.set(qn("w:val"), "164B73")
	properties.append(color)
	run.append(properties)
	text_element = OxmlElement("w:t")
	text_element.text = text
	run.append(text_element)
	hyperlink.append(run)
	paragraph._p.append(hyperlink)


def issue_brief_docx(brief: dict[str, Any]) -> bytes:
	"""Render all issue-brief page content as a source-linked Word document."""
	tag = str(brief.get("tag") or "")
	decision_count = int(brief.get("decision_count") or 0)
	semantics = brief.get("semantics") or {}
	title = f"Legal issue brief: {tag or '(empty tag)'}"
	document = Document(BytesIO(text_to_docx(title)))
	document.paragraphs[0].style = "Heading 1"
	document.add_paragraph(
		f"{decision_count} tagged decisions · exact active-taxonomy tag match"
	)
	if not tag:
		document.add_paragraph(
			"Enter a tag in category:value form, for example issue:procedural_fairness, "
			"then reload /issue-brief-ui?tag=…."
		)
	if not decision_count and tag:
		document.add_paragraph("No decisions are tagged with this issue.")

	if decision_count:
		document.add_heading("Decisions by year and outcome", level=2)
		years_table = document.add_table(rows=1, cols=3)
		years_table.style = "Table Grid"
		for cell, label in zip(
			years_table.rows[0].cells,
			("Year / outcome", "Decisions", "Unclassified / outcome share"),
		):
			cell.text = label
		for row in brief.get("years", []):
			year_cells = years_table.add_row().cells
			year_cells[0].text = str(
				row.get("year") if row.get("year") is not None else "Unknown"
			)
			year_cells[1].text = str(int(row.get("decision_count") or 0))
			year_cells[2].text = (
				f"Unclassified: {int(row.get('unclassified_count') or 0)}"
			)
			for outcome in row.get("outcome_splits", []):
				outcome_cells = years_table.add_row().cells
				outcome_cells[0].text = str(outcome.get("outcome") or "unclassified")
				outcome_cells[1].text = str(int(outcome.get("count") or 0))
				outcome_cells[2].text = (
					f"{float(outcome.get('percentage') or 0):.1f}% "
					f"(unclassified {int(outcome.get('unclassified_count') or 0)}; "
					f"denominator {int(outcome.get('denominator') or 0)})"
				)
		document.add_paragraph(
			"Each outcome percentage uses all tagged decisions in its year as the "
			"denominator, including unclassified decisions; unclassified count and "
			"denominator are shown beside every percentage."
		)

		document.add_heading("Courts", level=2)
		courts_table = document.add_table(rows=1, cols=2)
		courts_table.style = "Table Grid"
		courts_table.rows[0].cells[0].text = "Court"
		courts_table.rows[0].cells[1].text = "Decisions"
		for row in brief.get("courts", []):
			cells = courts_table.add_row().cells
			cells[0].text = str(row.get("court") or "Unspecified")
			cells[1].text = str(int(row.get("decision_count") or 0))

		document.add_heading("Top cited authorities", level=2)
		authorities = brief.get("top_authorities", [])
		if authorities:
			authority_table = document.add_table(rows=1, cols=3)
			authority_table.style = "Table Grid"
			for cell, label in zip(
				authority_table.rows[0].cells,
				("Authority", "Citation occurrences", "Tagged citing decisions"),
			):
				cell.text = label
			for row in authorities:
				cells = authority_table.add_row().cells
				paragraph = cells[0].paragraphs[0]
				label = str(row.get("citation") or row.get("title") or "Linked authority")
				url = str(row.get("url") or "")
				if url:
					_add_hyperlink(paragraph, label, url)
				else:
					paragraph.add_run(label)
				cells[1].text = str(int(row.get("citation_occurrences") or 0))
				cells[2].text = str(int(row.get("citing_decisions") or 0))
		else:
			document.add_paragraph("No resolved case citations found.")

		document.add_heading("Tagged decisions", level=2)
		decisions = brief.get("decisions", [])
		for row in decisions[:MAX_PRINT_DECISIONS]:
			paragraph = document.add_paragraph(style="List Bullet")
			label = str(row.get("citation") or row.get("title") or "Decision")
			url = str(row.get("url") or "")
			if url:
				_add_hyperlink(paragraph, label, url)
			else:
				paragraph.add_run(label)
			paragraph.add_run(f" — {row.get('court') or 'Unspecified'}")
			if row.get("outcome"):
				paragraph.add_run(f" · {row['outcome']}")
		if not decisions:
			document.add_paragraph("No linked decisions.")
		if len(decisions) or decision_count:
			shown_count = min(len(decisions), MAX_PRINT_DECISIONS)
			document.add_paragraph(
				f"Showing {shown_count} of {decision_count} decisions; "
				"see JSON brief for the complete list."
			)

	document.add_paragraph(f"Outcome source: {semantics.get('outcomes') or ''}")
	document.add_paragraph(f"Citation scope: {semantics.get('citations') or ''}")
	if semantics.get("tag_matching"):
		document.add_paragraph(f"Tag matching: {semantics['tag_matching']}")

	footer = document.sections[0].footer.paragraphs[0]
	footer.text = (
		"Generated from iLit data on "
		+ datetime.now(timezone.utc).date().isoformat()
	)
	for run in footer.runs:
		run.font.size = Pt(8)
	buffer = BytesIO()
	document.save(buffer)
	return buffer.getvalue()


def issue_brief_page_html(brief: dict[str, Any]) -> str:
	"""Render a compact printable view of an issue brief JSON payload."""
	tag = escape(str(brief.get("tag") or ""))
	decision_count = int(brief.get("decision_count") or 0)
	decisions = brief.get("decisions", [])
	semantics = brief.get("semantics") or {}

	if not decision_count:
		empty = '<p class="empty">No decisions are tagged with this issue.</p>'
		year_rows = court_rows = authority_rows = decision_rows = empty
	else:
		year_rows = "".join(
			"<tr>"
			f"<td>{escape(str(row.get('year') if row.get('year') is not None else 'Unknown'))}</td>"
			f"<td>{int(row.get('decision_count') or 0)}</td>"
			f"<td>Unclassified: {int(row.get('unclassified_count') or 0)}</td>"
			"</tr>"
			+ "".join(
				"<tr class=\"split\">"
				f"<td>{escape(str(outcome.get('outcome') or 'unclassified'))}</td>"
				f"<td>{int(outcome.get('count') or 0)}</td>"
				f"<td>{float(outcome.get('percentage') or 0):.1f}%"
				f" (unclassified {int(outcome.get('unclassified_count') or 0)};"
				f" denominator {int(outcome.get('denominator') or 0)})</td>"
				"</tr>"
				for outcome in row.get("outcome_splits", [])
			)
			for row in brief.get("years", [])
		)
		court_rows = "".join(
			f"<tr><td>{escape(str(row.get('court') or 'Unspecified'))}</td>"
			f"<td>{int(row.get('decision_count') or 0)}</td></tr>"
			for row in brief.get("courts", [])
		)
		authority_rows = "".join(
			"<tr>"
			f"<td><a href=\"{escape(str(row.get('url') or '#'), quote=True)}\">"
			f"{escape(str(row.get('citation') or row.get('title') or 'Linked authority'))}</a></td>"
			f"<td>{int(row.get('citation_occurrences') or 0)}</td>"
			f"<td>{int(row.get('citing_decisions') or 0)}</td></tr>"
			for row in brief.get("top_authorities", [])
		) or '<tr><td colspan="3">No resolved case citations found.</td></tr>'
		decision_parts = []
		for row in decisions[:MAX_PRINT_DECISIONS]:
			outcome = f" · {escape(str(row['outcome']))}" if row.get("outcome") else ""
			decision_parts.append(
				f"<li><a href=\"{escape(str(row.get('url') or '#'), quote=True)}\">"
				f"{escape(str(row.get('citation') or row.get('title') or 'Decision'))}</a>"
				f" — {escape(str(row.get('court') or 'Unspecified'))}{outcome}</li>"
			)
		decision_rows = "".join(decision_parts) or '<li>No linked decisions.</li>'
		decision_disclosure = (
			f'<p class="note decision-disclosure">Showing {len(decision_parts)} of '
			f'{decision_count} decisions; see JSON brief for the complete list.</p>'
		)

	empty_state = not decision_count
	return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Issue brief: {tag}</title>
<style>
*{{box-sizing:border-box}}body{{margin:0;background:#f2f1ec;color:#202522;font:13px/1.4 system-ui,sans-serif}}
main{{max-width:1000px;margin:24px auto;padding:28px;background:white;border:1px solid #d8d5ca}}
h1{{font-size:24px;margin:0 0 4px}}h2{{font-size:15px;margin:18px 0 6px;border-bottom:1px solid #d8d5ca;padding-bottom:4px}}
.meta,.note{{color:#535b56;font-size:11px}}.grid{{display:grid;grid-template-columns:1fr 1fr;gap:8px 20px}}
table{{width:100%;border-collapse:collapse;font-size:11px}}th,td{{padding:4px 6px;border-bottom:1px solid #e7e4da;text-align:left;vertical-align:top}}
th{{background:#f5f4ef}}.split td:first-child{{padding-left:18px}}ul{{margin:4px 0;padding-left:20px;columns:2}}li{{break-inside:avoid;margin-bottom:3px}}
a{{color:#164b73}}.empty{{padding:16px;border:1px dashed #aaa;color:#535b56}}
@media(max-width:650px){{main{{margin:0;padding:16px}}.grid{{grid-template-columns:1fr}}ul{{columns:1}}}}
@media print{{@page{{size:auto;margin:8mm}}body{{background:white;font-size:7.5pt;line-height:1.2}}main{{max-width:none;margin:0;padding:0;border:0}}
h1{{font-size:15pt;margin-bottom:2pt}}h2{{font-size:9pt;margin:5pt 0 2pt;padding-bottom:2pt}}
.grid{{grid-template-columns:1fr 1fr;gap:4pt 10pt}}table{{font-size:7pt}}th,td{{padding:1.5pt 3pt}}
ul{{columns:3;margin:2pt 0;padding-left:12pt}}li{{margin-bottom:1pt;break-inside:avoid}}
a{{color:inherit;text-decoration:none}}a[href]::after{{content:""}}.meta,.note{{font-size:6.5pt}}}}
</style></head>
<body><main>
<h1>Legal issue brief: {tag or "(empty tag)"}</h1>
<p class="meta">{decision_count} tagged decisions · exact active-taxonomy tag match</p>
{('<p class="empty">Enter a tag in category:value form, for example <code>issue:procedural_fairness</code>, then reload <code>/issue-brief-ui?tag=…</code>.</p>' if not tag else '')}
{('<p class="empty">No decisions are tagged with this issue.</p>' if empty_state and tag else '')}
{'' if empty_state else f'''<div class="grid">
<section><h2>Decisions by year and outcome</h2><table><thead><tr><th>Year / outcome</th><th>Decisions</th><th>Unclassified / outcome share</th></tr></thead><tbody>{year_rows}</tbody></table>
<p class="note">Each outcome percentage uses all tagged decisions in its year as the denominator, including unclassified decisions; unclassified count and denominator are printed beside every percentage.</p></section>
<section><h2>Courts</h2><table><thead><tr><th>Court</th><th>Decisions</th></tr></thead><tbody>{court_rows}</tbody></table>
<h2>Top cited authorities</h2><table><thead><tr><th>Authority</th><th>Citation occurrences</th><th>Tagged citing decisions</th></tr></thead><tbody>{authority_rows}</tbody></table></section>
</div><h2>Tagged decisions</h2>{decision_disclosure}<ul>{decision_rows}</ul>'''}
<p class="note">Outcome source: {escape(str(semantics.get("outcomes") or ""))}</p>
<p class="note">Citation scope: {escape(str(semantics.get("citations") or ""))}</p>
</main></body></html>"""
