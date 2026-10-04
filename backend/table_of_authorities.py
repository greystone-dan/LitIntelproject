"""Stateless parsing, local metadata resolution, and DOCX output for authorities."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from urllib.parse import urlsplit

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .citations import extract_case_citation_matches

MAX_INPUT_LINES = 200
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_NEUTRAL_CITATION_RE = re.compile(
    r"(?<![A-Za-z0-9])(?:\[)?((?:19|20)\d{2})(?:\])?\s+"
    r"([A-Za-z.]{1,12})\s+(\d{1,7})",
    re.IGNORECASE,
)
_REPORTED_CITATION_RE = re.compile(
    r"(?<![A-Za-z0-9])(?:\[((?:19|20)\d{2})\]|\(((?:19|20)\d{2})\)|"
    r"((?:19|20)\d{2}))\s+(\d{1,3})\s+([A-Za-z.]{1,12})\s+(\d{1,7})"
)
_PARAGRAPH_REF_RE = re.compile(
    r"\b(?:at\s+)?(?:paras?\.?|paragraphs?\.?)\s*((?:\d+[A-Za-z]?(?:\s*(?:[-–—,]|and)\s*"
    r"\d+[A-Za-z]?)*))",
    re.IGNORECASE,
)


class InputLineLimitError(ValueError):
    """Raised when a submission contains more than the permitted nonblank lines."""


@dataclass(frozen=True)
class CaseMetadata:
    case_id: int
    title: str
    court: str
    citation: str | None = None
    secondary_citation: str | None = None
    source_url: str | None = None


@dataclass
class Authority:
    citation: str
    title: str
    court: str | None
    paragraph_refs: list[str] = field(default_factory=list)
    canlii_url: str | None = None
    resolved: bool = False


@dataclass(frozen=True)
class ParsedCitation:
    citation: str
    paragraph_refs: tuple[str, ...]


def _citation_keys(value: str | None) -> set[str]:
    text = " ".join((value or "").upper().split())
    keys: set[str] = set()

    for match in _NEUTRAL_CITATION_RE.finditer(text):
        year, court, number = match.groups()
        court_key = re.sub(r"[^A-Z]", "", court)
        keys.add(f"{year}{court_key}{int(number)}")
        if court_key in {"FC", "FCT"}:
            keys.add(f"{year}FC{int(number)}")

    for match in _REPORTED_CITATION_RE.finditer(text):
        year = match.group(1) or match.group(2) or match.group(3)
        volume, reporter, page = match.group(4), match.group(5), match.group(6)
        keys.add(f"{year}{int(volume)}{re.sub(r'[^A-Z]', '', reporter)}{int(page)}")

    if not keys and text:
        compact = re.sub(r"[^A-Z0-9]", "", text)
        if compact:
            keys.add(compact)
    return keys


def _paragraph_refs(value: str | None) -> list[str]:
    if not value:
        return []
    refs: list[str] = []
    for match in _PARAGRAPH_REF_RE.finditer(value):
        refs.extend(
            re.findall(
                r"\d+[A-Za-z]?(?:\s*[-–—]\s*\d+[A-Za-z]?)?",
                match.group(1),
            )
        )
    refs = ["–".join(part.strip() for part in re.split(r"\s*[-—–]\s*", ref)) for ref in refs]
    return list(dict.fromkeys(refs))


def parse_submission(text: str) -> list[ParsedCitation]:
    """Extract case citations and attached paragraph references from pasted lines."""
    lines = [line.strip() for line in re.split(r"\r\n|\r|\n", text or "") if line.strip()]
    if len(lines) > MAX_INPUT_LINES:
        raise InputLineLimitError(
            f"Please limit the submission to {MAX_INPUT_LINES} nonblank lines "
            f"(received {len(lines)})."
        )

    parsed: list[ParsedCitation] = []
    for line in lines:
        if _positive_case_id(line) is not None:
            continue
        matches = extract_case_citation_matches(line)
        for index, match in enumerate(matches):
            next_start = matches[index + 1].offset_start if index + 1 < len(matches) else len(line)
            tail = line[match.offset_end:next_start]
            refs = _paragraph_refs(match.pinpoint) + _paragraph_refs(tail)
            citation = (match.normalized_citation or match.citation_text).strip(" ,;.")
            if citation:
                parsed.append(ParsedCitation(citation, tuple(dict.fromkeys(refs))))
    return parsed


def _positive_case_id(line: str) -> int | None:
    if not re.fullmatch(r"\d+", line):
        return None
    case_id = int(line)
    return case_id if case_id > 0 else None


def case_id_lookup_values(text: str) -> list[int]:
    """Return distinct positive, numeric-only lines as local case IDs."""
    return sorted(
        {
            case_id
            for line in re.split(r"\r\n|\r|\n", text or "")
            if (case_id := _positive_case_id(line.strip())) is not None
        }
    )


def citation_lookup_values(citations: list[ParsedCitation]) -> list[str]:
    """Return exact neutral/reported citation values for a selective metadata query."""
    values: set[str] = set()
    for citation in citations:
        for match in _NEUTRAL_CITATION_RE.finditer(citation.citation):
            year, court, number = match.groups()
            court_key = re.sub(r"[^A-Za-z]", "", court).upper()
            values.add(f"{year} {court_key} {int(number)}")
            if court_key == "FC":
                values.add(f"{year} FCT {int(number)}")
            elif court_key == "FCT":
                values.add(f"{year} FC {int(number)}")
        for match in _REPORTED_CITATION_RE.finditer(citation.citation):
            year = match.group(1) or match.group(2) or match.group(3)
            volume, reporter, page = match.group(4), match.group(5), match.group(6)
            values.add(
                f"{year} {int(volume)} {re.sub(r'[^A-Za-z0-9]', '', reporter).upper()} {int(page)}"
            )
            values.add(f"{year} {int(volume)} {reporter.upper()} {int(page)}")
    return sorted(values)


def _trusted_canlii_url(value: str | None) -> str | None:
    if not value:
        return None
    try:
        parsed = urlsplit(value)
    except ValueError:
        return None
    if (
        parsed.scheme != "https"
        or parsed.hostname != "www.canlii.org"
        or parsed.netloc.lower() != "www.canlii.org"
        or parsed.username
        or parsed.password
        or not re.fullmatch(
            r"/en/[^/]+(?:/[^/]+)?/doc/\d{4}/[A-Za-z0-9._-]+/[A-Za-z0-9._-]+\.html",
            parsed.path,
        )
    ):
        return None
    return value


def resolve_authorities(
    citations: list[ParsedCitation],
    local_cases: list[CaseMetadata],
) -> tuple[list[tuple[str, list[Authority]]], list[Authority]]:
    """Resolve only against supplied local case metadata; never performs network I/O."""
    index: dict[str, CaseMetadata | None] = {}
    for case in local_cases:
        for value in (case.citation, case.secondary_citation):
            for key in _citation_keys(value):
                if key not in index:
                    index[key] = case
                elif index[key] is not None and index[key].case_id != case.case_id:
                    index[key] = None

    resolved: dict[int, Authority] = {}
    unresolved: dict[str, Authority] = {}
    for parsed in citations:
        candidates = _citation_keys(parsed.citation)
        case = next((index[key] for key in sorted(candidates) if index.get(key) is not None), None)
        if case:
            authority = resolved.setdefault(
                case.case_id,
                Authority(
                    citation=case.citation or parsed.citation,
                    title=case.title,
                    court=case.court or "Court not recorded",
                    canlii_url=_trusted_canlii_url(case.source_url),
                    resolved=True,
                ),
            )
            authority.paragraph_refs.extend(parsed.paragraph_refs)
            authority.paragraph_refs[:] = list(dict.fromkeys(authority.paragraph_refs))
        else:
            key = min(candidates) if candidates else parsed.citation.upper()
            authority = unresolved.setdefault(
                key,
                Authority(
                    citation=parsed.citation,
                    title=parsed.citation,
                    court=None,
                ),
            )
            authority.paragraph_refs.extend(parsed.paragraph_refs)
            authority.paragraph_refs[:] = list(dict.fromkeys(authority.paragraph_refs))

    groups: dict[str, list[Authority]] = {}
    for authority in resolved.values():
        groups.setdefault(authority.court or "Court not recorded", []).append(authority)
    ordered_groups = [
        (court, sorted(rows, key=lambda row: (row.title.casefold(), row.citation.casefold())))
        for court, rows in sorted(groups.items(), key=lambda item: court_order_key(item[0]))
    ]
    return ordered_groups, sorted(
        unresolved.values(), key=lambda row: (row.title.casefold(), row.citation.casefold())
    )


def resolve_case_ids(
    case_ids: list[int],
    local_cases: list[CaseMetadata],
) -> tuple[list[tuple[str, list[Authority]]], list[Authority]]:
    """Resolve explicitly supplied local case IDs without citation parsing."""
    cases_by_id = {case.case_id: case for case in local_cases}
    resolved: dict[int, Authority] = {}
    not_found: list[Authority] = []
    for case_id in dict.fromkeys(case_ids):
        case = cases_by_id.get(case_id)
        if case is None:
            not_found.append(
                Authority(
                    citation=f"Case ID {case_id}",
                    title=f"Case ID {case_id}",
                    court=None,
                )
            )
            continue
        resolved[case_id] = Authority(
            citation=case.citation or f"Case ID {case_id}",
            title=case.title,
            court=case.court or "Court not recorded",
            canlii_url=_trusted_canlii_url(case.source_url),
            resolved=True,
        )

    groups: dict[str, list[Authority]] = {}
    for authority in resolved.values():
        groups.setdefault(authority.court or "Court not recorded", []).append(authority)
    ordered_groups = [
        (court, sorted(rows, key=lambda row: (row.title.casefold(), row.citation.casefold())))
        for court, rows in sorted(groups.items(), key=lambda item: court_order_key(item[0]))
    ]
    return ordered_groups, not_found


def combine_authority_results(
    *results: tuple[list[tuple[str, list[Authority]]], list[Authority]],
) -> tuple[list[tuple[str, list[Authority]]], list[Authority]]:
    """Combine citation and case-ID results without repeating authorities."""
    groups: dict[str, dict[tuple[str, str], Authority]] = {}
    not_found: dict[str, Authority] = {}
    for resolved_groups, unresolved in results:
        for court, authorities in resolved_groups:
            for authority in authorities:
                key = (authority.title.casefold(), authority.citation.casefold())
                combined = groups.setdefault(court, {}).setdefault(key, authority)
                combined.paragraph_refs[:] = list(
                    dict.fromkeys(combined.paragraph_refs + authority.paragraph_refs)
                )
        for authority in unresolved:
            key = authority.citation.casefold()
            combined = not_found.setdefault(key, authority)
            combined.paragraph_refs[:] = list(
                dict.fromkeys(combined.paragraph_refs + authority.paragraph_refs)
            )

    ordered_groups = [
        (
            court,
            sorted(
                rows.values(),
                key=lambda row: (row.title.casefold(), row.citation.casefold()),
            ),
        )
        for court, rows in sorted(groups.items(), key=lambda item: court_order_key(item[0]))
    ]
    return ordered_groups, sorted(
        not_found.values(), key=lambda row: (row.title.casefold(), row.citation.casefold())
    )


def court_order_key(court: str) -> tuple[int, str]:
    """Stable court precedence: SCC, federal appeal, federal trial, provincial appeal,
    provincial superior, other provincial/territorial courts, then other courts.
    Unlisted courts within a tier are ordered alphabetically.
    """
    normalized = re.sub(r"[^a-z]", " ", court.casefold())
    normalized = " ".join(normalized.split())
    compact = re.sub(r"[^a-z]", "", normalized)
    if "supreme court of canada" in normalized or compact in {"scc", "supremecourt"}:
        tier = 0
    elif "federal court of appeal" in normalized or compact in {"fca", "fctca"}:
        tier = 1
    elif "federal court" in normalized or compact in {"fc", "fct"}:
        tier = 2
    elif "court of appeal" in normalized or normalized.endswith(" appellate court"):
        tier = 3
    elif (
        "superior court" in normalized
        or "king bench" in normalized
        or "king s bench" in normalized
        or "queen bench" in normalized
        or "queen s bench" in normalized
    ):
        tier = 4
    elif "court" in normalized or "tribunal" in normalized:
        tier = 5
    else:
        tier = 6
    return tier, normalized


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def build_docx(
    groups: list[tuple[str, list[Authority]]],
    not_found: list[Authority],
) -> bytes:
    document = Document()
    document.add_heading("Table of Authorities", level=0)
    for court, authorities in groups:
        document.add_heading(court, level=1)
        for authority in authorities:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(authority.title)
            paragraph.add_run(f", {authority.citation}")
            if authority.paragraph_refs:
                paragraph.add_run(f", paras {', '.join(authority.paragraph_refs)}")
            if authority.canlii_url:
                paragraph.add_run(" ")
                _add_hyperlink(paragraph, "CanLII", authority.canlii_url)

    document.add_heading("Not found in local case metadata", level=1)
    if not not_found:
        document.add_paragraph("No unresolved authorities.")
    else:
        for authority in not_found:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(authority.citation)
            if authority.paragraph_refs:
                paragraph.add_run(f", paras {', '.join(authority.paragraph_refs)}")

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def build_table_of_authorities(text: str, local_cases: list[CaseMetadata]) -> bytes:
    parsed = parse_submission(text)
    groups, not_found = resolve_authorities(parsed, local_cases)
    return build_docx(groups, not_found)
