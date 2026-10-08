"""Build, and later score, a human review spreadsheet for paragraph roles. Offline.

build: picks ~N paragraphs, mostly where the runs disagree on role, writes an .xlsx with the paragraph text, the candidate
roles (run names hidden and order shuffled so the reviewer is not anchored) and a dropdown for the reviewer's own answer.
score: reads the filled sheet and reports, per run, how often its role equals the reviewer's answer.
"""

from __future__ import annotations

import argparse
import json
import random
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

ROLES = ["procedural_history", "background_facts", "issue_framing", "legal_test", "party_position",
	"evidence_assessment", "reasoning_application", "disposition", "other"]
HELP = [
	("procedural_history", "Steps and decisions before this proceeding (earlier hearings, appeals, who decided what), incl. the description of the decision under review."),
	("background_facts", "The claimant's or applicant's own life events and circumstances."),
	("issue_framing", "What the court will or will not decide, list of issues, standard of review to apply, overall approach."),
	("legal_test", "Governing law: statutes, rules, tests, what cases held."),
	("party_position", "Reports what a party or the decision-maker below argued, found or reasoned (describing, not deciding)."),
	("evidence_assessment", "Mainly summarises or weighs a specific item of evidence or the record, without a conclusion yet."),
	("reasoning_application", "The court's own analysis or conclusion on a point."),
	("disposition", "The court's result or order (allowed, dismissed, remitted, costs), incl. in an opening or closing summary."),
	("other", "Headings, style of cause, counsel lists, page furniture."),
]


def load(directory: Path) -> dict[tuple[int, int], str]:
	out = {}
	for path in directory.glob("case_*_tags.json"):
		data = json.loads(path.read_text(encoding="utf-8"))
		for a in data["assessments"]:
			out[(int(data["case_id"]), int(a["paragraph_index"]))] = a["role"]
	return out


def build(args) -> None:
	from openpyxl import Workbook
	from openpyxl.worksheet.datavalidation import DataValidation
	from openpyxl.styles import Alignment, Font
	from compare_tags import paragraph_texts

	names = [d.name for d in args.runs]
	runs = [load(d) for d in args.runs]
	reference = json.loads((Path(__file__).parent / "role_reference_60.json").read_text(encoding="utf-8"))["roles"]
	skip = {(int(c), int(i)) for c, d in reference.items() for i in d}
	common = set(runs[0])
	for r in runs[1:]:
		common &= set(r)
	common -= skip
	disagree = sorted(k for k in common if len({r[k] for r in runs}) > 1)
	agree = sorted(k for k in common if len({r[k] for r in runs}) == 1)
	rng = random.Random(args.seed)
	n_dis = min(len(disagree), int(args.n * 0.8))
	picked = rng.sample(disagree, n_dis) + rng.sample(agree, min(len(agree), args.n - n_dis))
	rng.shuffle(picked)
	wb = Workbook()
	ws = wb.active
	ws.title = "Review"
	ws.append(["#", "Case", "Para", "Paragraph text (first 700 characters)", "Candidate role 1", "Candidate role 2", "Candidate role 3", "YOUR ANSWER (pick the best role; may differ from candidates)"])
	key = []
	texts: dict[int, dict[int, str]] = {}
	for number, (case_id, index) in enumerate(picked, 1):
		texts.setdefault(case_id, paragraph_texts(case_id))
		cands = sorted({r[(case_id, index)] for r in runs})
		rng.shuffle(cands)
		ws.append([number, case_id, index, texts[case_id].get(index, "")[:700], *(cands + [""] * (3 - len(cands)))[:3], ""])
		key.append({"number": number, "case_id": case_id, "paragraph": index, "runs": {n: r[(case_id, index)] for n, r in zip(names, runs)}})
	for col, width in zip("ABCDEFGH", (5, 8, 6, 90, 22, 22, 22, 30)):
		ws.column_dimensions[col].width = width
	for row in ws.iter_rows(min_row=2):
		for cell in row:
			cell.alignment = Alignment(wrap_text=True, vertical="top")
	for cell in ws[1]:
		cell.font = Font(bold=True)
		cell.alignment = Alignment(wrap_text=True)
	ws.freeze_panes = "E2"
	dv = DataValidation(type="list", formula1='"' + ",".join(ROLES) + '"', allow_blank=True)
	ws.add_data_validation(dv)
	dv.add(f"H2:H{len(picked) + 1}")
	guide = wb.create_sheet("Role guide")
	guide.append(["Role", "Meaning"])
	for row in HELP:
		guide.append(list(row))
	guide.column_dimensions["A"].width = 24
	guide.column_dimensions["B"].width = 110
	wb.save(args.out)
	Path(f"{args.out}.key.json").write_text(json.dumps({"runs": names, "items": key}, indent=1), encoding="utf-8")
	print(json.dumps({"rows": len(picked), "disagreements": n_dis, "xlsx": str(args.out)}))


def score(args) -> None:
	from openpyxl import load_workbook

	key = json.loads(Path(f"{args.xlsx}.key.json").read_text(encoding="utf-8"))
	ws = load_workbook(args.xlsx)["Review"]
	answers = {int(r[0].value): (r[7].value or "").strip() for r in ws.iter_rows(min_row=2) if r[0].value}
	totals = {n: [0, 0] for n in key["runs"]}
	for item in key["items"]:
		answer = answers.get(item["number"])
		if answer:
			for name, role in item["runs"].items():
				totals[name][1] += 1
				totals[name][0] += role == answer
	print(json.dumps({n: {"right": r, "of": t, "accuracy": round(r / t, 3) if t else None} for n, (r, t) in totals.items()}, indent=1))
	print("note: the sheet over-samples disagreements, so these are NOT overall accuracy; see the report for the adjustment")


def main() -> int:
	parser = argparse.ArgumentParser(description=__doc__)
	sub = parser.add_subparsers(dest="cmd", required=True)
	b = sub.add_parser("build")
	b.add_argument("runs", nargs="+", type=Path)
	b.add_argument("--out", type=Path, required=True)
	b.add_argument("--n", type=int, default=100)
	b.add_argument("--seed", type=int, default=5)
	s = sub.add_parser("score")
	s.add_argument("xlsx", type=Path)
	args = parser.parse_args()
	(build if args.cmd == "build" else score)(args)
	return 0


if __name__ == "__main__":
	raise SystemExit(main())
