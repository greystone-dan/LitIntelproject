import json
import re
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from backend.deidentify import deidentify_text, reidentify_text, text_from_upload
from backend.main import app
from backend import resource_limits

SAMPLE = """AFFIDAVIT OF MARIA ELENA LOPEZ
Court File No. IMM-4521-25   RPD File No. TB9-14699
I, Maria Elena Lopez, born 12 March 1988 in Oaxaca, Mexico, UCI 1234-5678, make oath.
1. Ms. Lopez worked as a nurse at Hospital General. Her husband, Juan Perez, is 41 years old.
2. Application number W305123456. Passport No. G12345678. SIN 046 454 286.
3. I live at Apt 4, 123 Main Street, Toronto, ON M5V 2T6. Tel. (416) 555-0199, email maria.lopez@example.com.
4. On May 3, 2019, PEREZ was threatened. During 2018-2019 we hid. Client ID number 87654321.
5. See Vavilov, 2019 SCC 65 at para 23."""


def _run(**kwargs):
	return deidentify_text(SAMPLE, names=["Maria Elena Lopez", "Juan Perez"], details=["Hospital General"], **kwargs)


def test_identifying_details_are_removed():
	result = _run()
	text = result["text"]
	for secret in [
		"Maria", "Lopez", "LOPEZ", "Perez", "PEREZ", "IMM-4521-25", "TB9-14699", "12 March 1988",
		"1234-5678", "Hospital General", "41 years old", "W305123456", "G12345678", "046 454 286",
		"123 Main Street", "M5V 2T6", "555-0199", "maria.lopez@example.com", "May 3, 2019", "87654321",
	]:
		assert secret not in text, secret
	assert result["warnings"] == []


def test_country_employment_and_legal_citations_are_kept():
	text = _run()["text"]
	for kept in ["Oaxaca, Mexico", "nurse", "2018-2019", "Vavilov, 2019 SCC 65", "Toronto"]:
		assert kept in text, kept


def test_same_person_gets_same_placeholder_and_name_parts_are_linked():
	text = _run()["text"]
	assert text.count("[PERSON_1]") == 1
	assert "AFFIDAVIT OF [PERSON_1_CAPS]" in text
	assert "Ms. [PERSON_1_SURNAME]" in text
	assert "[PERSON_2_SURNAME_CAPS] was threatened" in text


def test_round_trip_restores_exact_text_even_if_placeholders_are_reformatted():
	result = _run()
	processed = result["text"].replace("[PERSON_1_SURNAME]", "[Person 1 surname]").replace("[UCI_1]", "[uci-1]")
	restored = reidentify_text(processed, result["key"])
	assert restored["text"] == SAMPLE
	assert restored["warnings"] == []


def test_restore_flags_unknown_placeholders_and_rejects_bad_keys():
	key = _run()["key"]
	result = reidentify_text("[PERSON_9] met [PERSON_1].", key)
	assert result["text"] == "[PERSON_9] met Maria Elena Lopez."
	assert "[PERSON_9]" in result["warnings"][0]
	with pytest.raises(ValueError):
		reidentify_text("x", {"entries": {}})


def test_missing_name_is_warned_and_categories_can_be_switched_off():
	result = deidentify_text("Born 12 March 1988.", names=["Nobody Here"], categories=["UCI"])
	assert "12 March 1988" in result["text"]
	assert any("Nobody Here" in warning for warning in result["warnings"])


def test_docx_extraction_includes_tables_and_headers():
	document = Document()
	document.sections[0].header.paragraphs[0].text = "Client: Maria Lopez"
	document.add_paragraph("Body text")
	table = document.add_table(rows=1, cols=2)
	table.cell(0, 0).text = "UCI"
	table.cell(0, 1).text = "1234-5678"
	buffer = BytesIO()
	document.save(buffer)
	text = text_from_upload("form.docx", buffer.getvalue())
	assert "Client: Maria Lopez" in text and "UCI | 1234-5678" in text and "Body text" in text


def test_unsupported_upload_is_rejected():
	with pytest.raises(ValueError):
		text_from_upload("photo.jpg", b"data")


def test_api_round_trip_without_database():
	client = TestClient(app)
	assert client.get("/deidentify").status_code == 200
	response = client.post("/api/deidentify", data={"text": SAMPLE, "names": "Maria Elena Lopez\nJuan Perez"})
	assert response.status_code == 200
	assert response.headers["cache-control"] == "no-store"
	body = response.json()
	assert "Lopez" not in body["text"]
	restored = client.post("/api/reidentify", data={"text": body["text"], "key": json.dumps(body["key"])})
	assert restored.status_code == 200
	assert "Juan Perez" in restored.json()["text"]
	bad = client.post("/api/reidentify", data={"text": "x", "key": "not json"})
	assert bad.status_code == 422
	docx = client.post("/api/deidentify/docx", data={"text": "Hello\n\nWorld", "filename": "a b.docx"})
	assert docx.status_code == 200 and docx.content[:2] == b"PK"


def test_deidentify_upload_limit_returns_http_413(monkeypatch):
	monkeypatch.setattr(resource_limits, "MAX_UPLOAD_BYTES", 16)
	response = TestClient(app).post(
		"/api/deidentify",
		files={"file": ("brief.txt", b"x" * 17, "text/plain")},
	)

	assert response.status_code == 413
	assert "uploaded file exceeds" in response.json()["detail"]


def test_deidentify_pasted_text_limit_returns_http_413(monkeypatch):
	monkeypatch.setattr(resource_limits, "MAX_PASTED_TEXT_CHARS", 8)
	response = TestClient(app).post("/api/deidentify", data={"text": "123456789"})

	assert response.status_code == 413
	assert response.json()["detail"] == "The pasted text exceeds the 8 character limit."


def test_deidentify_docx_output_text_limit_returns_http_413(monkeypatch):
	monkeypatch.setattr(resource_limits, "MAX_PASTED_TEXT_CHARS", 8)
	response = TestClient(app).post(
		"/api/deidentify/docx",
		data={"text": "123456789"},
	)

	assert response.status_code == 413
	assert response.json()["detail"] == "The pasted text exceeds the 8 character limit."


# ---- Automatic name detection ----------------------------------------------

spacy = pytest.importorskip("spacy")
try:
	spacy.load("en_core_web_md")
	HAS_MODEL = True
except OSError:
	HAS_MODEL = False
needs_model = pytest.mark.skipif(not HAS_MODEL, reason="en_core_web_md is not installed")

NARRATIVE = """BASIS OF CLAIM

My name is Oluwaseun Adebayo. I am a citizen of Nigeria and I worked as a teacher.
1. In 2016 I married my wife, Folake Adebayo. Her uncle, Tunde Bakare, accused me of being gay.
2. Mr. Bakare told the elders. One attacker, Ibrahim Musa, said they would kill me if I stayed in Lagos.
3. The Adebayos fled. Prices rose and my brother Harjinder Singh Sandhu sent money.
4. The decision must be reasonable: Canada (Minister of Citizenship and Immigration) v Vavilov, 2019 SCC 65.
   As Justice Grammond said, Vavilov applies.
"""


@needs_model
def test_names_are_found_without_being_typed():
	result = deidentify_text(NARRATIVE, auto_names=True)
	text = result["text"]
	for name in ["Oluwaseun", "Adebayo", "Folake", "Tunde", "Bakare", "Ibrahim", "Musa", "Harjinder", "Sandhu"]:
		assert name not in text, name
	assert "The [PERSON" in text and "Adebayos" not in text  # plural surname
	assert "Prices rose" in text  # ordinary words that are also names survive
	assert "Nigeria" in text and "Lagos" in text and "teacher" in text


@needs_model
def test_cited_cases_and_judges_are_kept_but_can_be_forced():
	result = deidentify_text(NARRATIVE, auto_names=True)
	assert "v Vavilov, 2019 SCC 65" in result["text"] and "Justice Grammond" in result["text"]
	reasons = {k["name"]: k["reason"] for k in result["kept_names"]}
	assert reasons["Vavilov"] == "Party in a cited case"
	forced = deidentify_text(NARRATIVE, names=["Vavilov"], auto_names=True)
	assert "Vavilov" not in forced["text"]


@needs_model
def test_never_hide_list_overrides_detection():
	result = deidentify_text(NARRATIVE, auto_names=True, never_hide=["Ibrahim Musa"])
	assert "Ibrahim Musa" in result["text"]
	assert any(k["name"] == "Ibrahim Musa" for k in result["kept_names"])


@needs_model
def test_court_style_of_cause_names_are_hidden_and_restore_exactly():
	text = (
		"BETWEEN:\nSeonhee lee a.k.a. Younglan LEE, Seonwoo JUNG et al\nApplicants\nand\n"
		"THE MINISTER OF CITIZENSHIP AND IMMIGRATION\nRespondent\n"
		"[1] Ms. Lee and Mr. Jung seek judicial review. JUNG, Seonwoo was born in Seoul.\n"
	)
	result = deidentify_text(text, auto_names=True)
	for name in ["Seonhee", "Younglan", "LEE", "Lee", "Seonwoo", "JUNG", "Jung"]:
		assert not re.search(rf"\b{name}\b", result["text"]), name
	assert "Seoul" in result["text"]
	assert reidentify_text(result["text"], result["key"])["text"] == text


def test_auto_names_off_hides_only_typed_names():
	result = deidentify_text("Juan Perez met Ana Diaz.", names=["Juan Perez"], auto_names=False)
	assert result["text"] == "[PERSON_1] met Ana Diaz." and result["detected_names"] == []
