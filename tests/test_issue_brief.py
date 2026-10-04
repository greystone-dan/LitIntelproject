from datetime import date
from io import BytesIO
import re

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pgvector.sqlalchemy import Vector
from sqlalchemy import create_engine
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from backend.analytics_service import ACTIVE_TAG_TAXONOMY_VERSION
from backend.database import Base, Case, CaseTag, Citation, get_db
from backend.main import app
from backend.pages.issue_brief import issue_brief_page_html


@compiles(Vector, "sqlite")
def _sqlite_vector_type(_type, _compiler, **_kw):
    return "JSON"


@pytest.fixture
def issue_brief_client():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(
        engine,
        tables=[Case.__table__, CaseTag.__table__, Citation.__table__],
    )
    session = Session(engine)
    cases = [
        Case(
            id=1,
            title="Authority v. Canada",
            court="Federal Court",
            date=date(2020, 1, 1),
            citation="2020 FC 1",
            source_url="https://example.test/authority",
        ),
        Case(
            id=2,
            title="Applicant v. Canada",
            court="Federal Court",
            date=date(2024, 2, 1),
            citation="2024 FC 2",
            metadata_json={"reader_extracted": {"decision outcome": "allowed"}},
        ),
        Case(
            id=3,
            title="Another Applicant v. Canada",
            court="Federal Court of Appeal",
            date=date(2024, 4, 1),
            citation="2024 FCA 3",
            metadata_json={"reader_extracted": {}},
        ),
        Case(
            id=4,
            title="Third Applicant v. Canada",
            court="Federal Court",
            date=date(2025, 3, 1),
            citation="2025 FC 4",
            metadata_json={"reader_extracted": {"decision outcome": "dismissed"}},
        ),
    ]
    session.add_all(cases)
    session.flush()
    for case_id in (2, 3, 4):
        session.add(
            CaseTag(
                case_id=case_id,
                category="issue",
                value="fairness",
                score=1.0,
                evidence="procedural fairness",
                source="test",
                taxonomy_version=ACTIVE_TAG_TAXONOMY_VERSION,
            )
        )
    session.add_all(
        [
            Citation(source_case_id=2, target_case_id=1, citation_kind="case", citation_text="2020 FC 1"),
            Citation(source_case_id=2, target_case_id=1, citation_kind="case", citation_text="Authority"),
            Citation(source_case_id=3, target_case_id=1, citation_kind="case", citation_text="2020 FC 1"),
            Citation(source_case_id=3, target_case_id=None, citation_kind="case", citation_text="Unresolved FC case"),
        ]
    )
    session.commit()

    def override_get_db():
        yield session

    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app)
    try:
        yield client
    finally:
        app.dependency_overrides.pop(get_db, None)
        session.close()
        engine.dispose()


def test_issue_brief_json_contract(issue_brief_client):
    response = issue_brief_client.get("/issue-brief", params={"tag": "issue:fairness"})

    assert response.status_code == 200
    brief = response.json()
    assert brief["decision_count"] == 3
    assert [row["year"] for row in brief["years"]] == [2024, 2025]
    splits = [split for year in brief["years"] for split in year["outcome_splits"]]
    assert {split["outcome"] for split in splits} == {"allowed", "dismissed", "unclassified"}
    for year in brief["years"]:
        for split in year["outcome_splits"]:
            assert split["denominator"] == year["decision_count"]
            assert split["unclassified_count"] == year["unclassified_count"]
            assert "percentage" in split
    outcomes_2024 = {row["outcome"]: row["percentage"] for row in brief["years"][0]["outcome_splits"]}
    assert outcomes_2024 == {"allowed": 50.0, "unclassified": 50.0}
    assert brief["years"][1]["outcome_splits"][0]["percentage"] == 100.0
    assert brief["years"][0]["unclassified_count"] == 1
    authority = brief["top_authorities"][0]
    assert authority["citation_occurrences"] == 3
    assert authority["citing_decisions"] == 2
    assert authority["url"] == "/case-reader?case_id=1"
    assert len(brief["decisions"]) == 3
    assert all(row["url"].startswith("/case-reader?case_id=") for row in brief["decisions"])
    assert {row["court"] for row in brief["courts"]} == {"Federal Court", "Federal Court of Appeal"}


def test_issue_brief_printable_ui(issue_brief_client):
    response = issue_brief_client.get("/issue-brief-ui", params={"tag": "issue:fairness"})

    assert response.status_code == 200
    assert "@media print" in response.text
    assert "Decisions by year and outcome" in response.text
    assert "unclassified 1; denominator 2" in response.text
    assert 'href="/case-reader?case_id=2"' in response.text
    assert "Top cited authorities" in response.text


def test_issue_brief_print_ui_bounds_decision_list():
    decisions = [
        {
            "citation": f"2024 FC {number}",
            "court": "Federal Court",
            "url": f"/case-reader?case_id={number}",
        }
        for number in range(1, 15)
    ]
    brief = {
        "tag": "issue:fairness",
        "decision_count": len(decisions),
        "decisions": decisions,
        "years": [],
        "courts": [],
        "top_authorities": [],
        "semantics": {},
    }

    html = issue_brief_page_html(brief)

    assert html.count("<li><a href=") == 12
    assert "2024 FC 12" in html
    assert "2024 FC 13" not in html
    assert "2024 FC 14" not in html
    assert "Showing 12 of 14 decisions; see JSON brief for the complete list." in html
    assert "@media print" in html
    assert "@page{size:auto;margin:8mm}" in html
    assert "ul{columns:3" in html


def test_issue_brief_empty_tag(issue_brief_client):
    response = issue_brief_client.get("/issue-brief", params={"tag": ""})

    assert response.status_code == 200
    assert response.json()["decision_count"] == 0
    assert response.json()["years"] == []
    assert response.json()["top_authorities"] == []

    ui = issue_brief_client.get("/issue-brief-ui", params={"tag": ""})
    assert ui.status_code == 200
    assert "Enter a tag in category:value form" in ui.text


def test_issue_brief_docx_route_content_links_and_headers(issue_brief_client):
    response = issue_brief_client.get(
        "/issue-brief.docx", params={"tag": "issue:fairness"}
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert response.headers["content-disposition"] == (
        'attachment; filename="issue-brief-issue-fairness.docx"'
    )
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["pragma"] == "no-cache"

    document = Document(BytesIO(response.content))
    text = "\n".join(document._element.xpath(".//w:t/text()"))
    assert "Legal issue brief: issue:fairness" in text
    assert "3 tagged decisions" in text
    assert "Decisions by year and outcome" in text
    assert "denominator 2" in text
    assert "Federal Court of Appeal" in text
    assert "Top cited authorities" in text
    assert "2020 FC 1" in text
    assert "Citation occurrences" in text
    assert "Tagged decisions" in text
    assert "2024 FC 2" in text
    assert "Outcome source:" in text
    assert "Citation scope:" in text
    assert "Tag matching:" in text
    footer = document.sections[0].footer.paragraphs[0].text
    assert re.fullmatch(r"Generated from iLit data on \d{4}-\d{2}-\d{2}", footer)
    targets = {rel.target_ref for rel in document.part.rels.values() if rel.is_external}
    assert "/case-reader?case_id=1" in targets
    assert "/case-reader?case_id=2" in targets


def test_issue_brief_docx_empty_state_and_tag_bounds(issue_brief_client):
    response = issue_brief_client.get("/issue-brief.docx", params={"tag": ""})
    assert response.status_code == 200
    document = Document(BytesIO(response.content))
    text = "\n".join(document._element.xpath(".//w:t/text()"))
    assert "0 tagged decisions" in text
    assert "Enter a tag in category:value form" in text

    oversized_tag = issue_brief_client.get(
        "/issue-brief.docx", params={"tag": "x" * 357}
    )
    assert oversized_tag.status_code == 422
