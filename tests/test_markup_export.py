"""Word export with margin notes as comments: package validity, anchoring, escaping, route (SQLite, offline)."""

import io
import zipfile
from datetime import date
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pgvector.sqlalchemy import Vector
from sqlalchemy import create_engine
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from backend import routes
from backend.database import Base, Case
from backend.markup_export import Comment, build_markup_docx

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
TEXT = "Intro line\n\nREASONS\n\n[1] The decision was reasonable & fair <ok>.\n\n[2] Second paragraph cites Vavilov, 2019 SCC 65 at para 7."


@compiles(Vector, "sqlite")
def _sqlite_vector_type(_type, _compiler, **_kw):
    return "JSON"


def _starts():
    return {"p1": TEXT.index("[1]"), "p2": TEXT.index("[2]")}


def _parts(blob):
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        return {n: z.read(n) for n in z.namelist()}


def test_docx_is_valid_and_carries_comments_on_the_right_paragraphs():
    s = _starts()
    blob = build_markup_docx(
        title="Doe v. Canada <1>", subtitle="2024 FC 1 · FC", full_text=TEXT,
        comments=[
            Comment(s["p1"], "Outcome", "Allowed.\nVerified at para 1"),
            Comment(s["p2"], "Citation", "Vavilov: reasonableness", quote="2019 SCC 65 at para 7"),
            Comment(None, "Judge", "Justice X"),
            Comment(s["p2"], "My note", "check this", author="My note"),
        ],
        highlights=[s["p1"]],
    )
    parts = _parts(blob)
    for name in ("word/document.xml", "word/comments.xml", "[Content_Types].xml", "_rels/.rels"):
        ET.fromstring(parts[name])  # every part is well-formed XML
    comments = ET.fromstring(parts["word/comments.xml"]).findall(f"{W}comment")
    assert len(comments) == 4 and [c.get(f"{W}id") for c in comments] == ["0", "1", "2", "3"]
    assert comments[3].get(f"{W}author") == "My note"
    first = "".join(t.text for t in comments[0].iter(f"{W}t"))
    assert first.startswith("Outcome") and "Verified at para 1" in first
    doc = Document(io.BytesIO(blob))  # python-docx can open it
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "Doe v. Canada <1>" and any("reasonable & fair <ok>" in t for t in texts)
    root = ET.fromstring(parts["word/document.xml"])
    refs = [r.get(f"{W}id") for r in root.iter(f"{W}commentReference")]
    assert sorted(refs) == ["0", "1", "2", "3"]


def test_quote_anchors_inside_the_paragraph_and_missing_quote_covers_it_all():
    s = _starts()
    blob = build_markup_docx(
        title="t", subtitle="", full_text=TEXT,
        comments=[Comment(s["p2"], "a", "x", quote="2019 SCC 65"), Comment(s["p2"], "b", "y", quote="not in the text")],
    )
    root = ET.fromstring(_parts(blob)["word/document.xml"])
    para = next(p for p in root.iter(f"{W}p") if "Second paragraph" in "".join(t.text or "" for t in p.iter(f"{W}t")))
    kids = list(para)
    tags = [k.tag.replace(W, "") for k in kids]
    assert tags.count("commentRangeStart") == 2 and tags.count("commentRangeEnd") == 2
    inside = []
    open_ids = set()
    for k in kids:
        name = k.tag.replace(W, "")
        if name == "commentRangeStart":
            open_ids.add(k.get(f"{W}id"))
        elif name == "commentRangeEnd":
            open_ids.discard(k.get(f"{W}id"))
        elif name == "r" and "0" in open_ids and k.find(f"{W}t") is not None:
            inside.append(k.find(f"{W}t").text)
    assert "".join(inside) == "2019 SCC 65"  # the quoted words only


def test_highlight_and_heading_and_bad_characters():
    s = _starts()
    blob = build_markup_docx(
        title="t\x00x", subtitle="", full_text=TEXT,
        comments=[Comment(999999, "Lost", "block no longer exists"), Comment(s["p1"], "L\x0b", "t\x08")],
        highlights=[s["p1"]],
    )
    parts = _parts(blob)
    xml = parts["word/document.xml"].decode()
    assert 'w:highlight w:val="yellow"' in xml and "<w:b/>" in xml  # highlighted paragraph, bold heading
    assert "Notes not attached to a paragraph" in xml  # orphan comments are kept, not dropped
    ET.fromstring(parts["word/comments.xml"])


def test_comment_count_and_text_are_capped():
    many = [Comment(None, "l", "x" * 9000) for _ in range(3100)]
    parts = _parts(build_markup_docx(title="t", subtitle="", full_text=TEXT, comments=many))
    comments = ET.fromstring(parts["word/comments.xml"]).findall(f"{W}comment")
    assert len(comments) == 3000
    assert len("".join(t.text for t in comments[0].iter(f"{W}t"))) <= 4000 + 120


@pytest.fixture
def client():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine, tables=[Case.__table__])
    with Session(engine) as db:
        db.add(Case(id=7, title="Doe v. Canada", citation="2024 FC 1", court="FC", date=date(2024, 1, 2), full_text=TEXT))
        db.add(Case(id=8, title="Empty", court="FC", date=date(2024, 1, 2), full_text=""))
        db.commit()
        app = FastAPI()
        app.include_router(routes.router)
        app.dependency_overrides[routes.get_db] = lambda: db
        with TestClient(app) as test_client:
            yield test_client


def test_route_returns_docx_and_rejects_missing_or_oversize(client):
    ok = client.post("/cases/7/markup-export", json={"comments": [{"block": None, "label": "Judge", "text": "X"}]})
    assert ok.status_code == 200 and ok.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "ilit-case-7-markup.docx" in ok.headers["content-disposition"]
    assert Document(io.BytesIO(ok.content)).paragraphs[0].text == "Doe v. Canada"
    assert client.post("/cases/999/markup-export", json={"comments": []}).status_code == 404
    assert client.post("/cases/8/markup-export", json={"comments": []}).status_code == 404
    assert client.post("/cases/7/markup-export", json={"comments": [{"text": "x" * 4001}]}).status_code == 422
