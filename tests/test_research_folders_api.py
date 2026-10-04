import csv
import io
from datetime import date

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend import routes


class _Rows:
    def __init__(self, rows):
        self.rows = rows

    def mappings(self):
        return self

    def all(self):
        return self.rows


class _ReadOnlySession:
    def __init__(self, rows):
        self.rows = rows
        self.statements = []

    def execute(self, statement):
        self.statements.append(statement)
        return _Rows(self.rows)


def _client(session):
    app = FastAPI()
    app.include_router(routes.router)
    app.dependency_overrides[routes.get_db] = lambda: session
    return TestClient(app)


def _case(case_id, *, citation, title, court, case_date, decision=None, government=None):
    return {
        "id": case_id,
        "citation": citation,
        "title": title,
        "court": court,
        "date": case_date,
        "decision_outcome": decision,
        "government_outcome": government,
    }


def test_folder_manager_page_is_served_without_database_access():
    session = _ReadOnlySession([])

    response = _client(session).get("/research-folders")

    assert response.status_code == 200
    assert 'id="researchFoldersApp"' in response.text
    assert "ilit.researchFolders.v1" in response.text
    assert not session.statements


def test_export_csv_limits_output_and_protects_formula_cells():
    session = _ReadOnlySession(
        [
            _case(
                22,
                citation="=HYPERLINK(\"bad\")",
                title="Second case",
                court="Court",
                case_date=date(2024, 2, 3),
                decision="unknown",
            ),
            _case(
                11,
                citation="2024 FC 10",
                title="+SUM(A1:A2)",
                court="Court",
                case_date=date(2023, 1, 2),
                government="won",
            ),
        ]
    )

    response = _client(session).post(
        "/api/research-folders/export",
        json={
            "case_ids": [11, 22, 11],
            "notes": {"11": "@review", "22": "ordinary note"},
            "format": "csv",
        },
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert response.headers["content-disposition"] == (
        'attachment; filename="research-folder-export.csv"'
    )
    assert response.headers["cache-control"] == "no-store"
    rows = list(csv.reader(io.StringIO(response.text.lstrip("\ufeff"))))
    assert rows == [
        ["Citation", "Name", "Court", "Date", "Outcome", "Note"],
        ["2024 FC 10", "'+SUM(A1:A2)", "Court", "2023-01-02", "won", "'@review"],
        ["'=HYPERLINK(\"bad\")", "Second case", "Court", "2024-02-03", "unclassified", "ordinary note"],
    ]
    assert len(session.statements) == 1
    statement = str(session.statements[0]).lower()
    assert "full_text" not in statement
    assert "summary" not in statement


def test_export_docx_contains_only_six_requested_columns_and_selected_cases():
    session = _ReadOnlySession(
        [
            _case(
                5,
                citation=None,
                title="Example name",
                court="Federal Court",
                case_date=None,
                decision="dismissed",
            )
        ]
    )

    response = _client(session).post(
        "/api/research-folders/export",
        json={"case_ids": [5], "notes": {"5": "Read paragraph 4"}, "format": "docx"},
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    document = Document(io.BytesIO(response.content))
    assert len(document.tables) == 1
    assert [cell.text for cell in document.tables[0].rows[0].cells] == [
        "Citation",
        "Name",
        "Court",
        "Date",
        "Outcome",
        "Note",
    ]
    assert [cell.text for cell in document.tables[0].rows[1].cells] == [
        "",
        "Example name",
        "Federal Court",
        "",
        "dismissed",
        "Read paragraph 4",
    ]
    assert not document.paragraphs[0].text


def test_export_rejects_more_than_500_distinct_ids_before_database_access():
    session = _ReadOnlySession([])

    response = _client(session).post(
        "/api/research-folders/export",
        json={"case_ids": list(range(1, 502)), "format": "csv"},
    )

    assert response.status_code == 422
    assert not session.statements


def test_export_rejects_more_than_500_requested_ids_even_when_duplicated():
    session = _ReadOnlySession([])

    response = _client(session).post(
        "/api/research-folders/export",
        json={"case_ids": [1] * 501, "format": "csv"},
    )

    assert response.status_code == 422
    assert not session.statements


def test_export_rejects_notes_for_unselected_case():
    session = _ReadOnlySession([])

    response = _client(session).post(
        "/api/research-folders/export",
        json={"case_ids": [1], "notes": {"2": "not selected"}},
    )

    assert response.status_code == 422
    assert not session.statements
